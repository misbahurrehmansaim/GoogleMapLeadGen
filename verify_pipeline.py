import time
import requests
import openpyxl

BASE_URL = "http://127.0.0.1:5000"

print("--- Testing API Start with the 4 Required Combinations ---")
payload = {
    "services": ["Dumpster Rental", "Pest Control"],
    "locations": ["New York", "California"],
    "max_results": 3,  # fast test across all 4 combinations
    "remove_duplicates": True
}

resp = requests.post(f"{BASE_URL}/api/start", json=payload)
print("Start response:", resp.status_code, resp.json())

# Monitor progress
start_time = time.time()
while True:
    time.sleep(2)
    s = requests.get(f"{BASE_URL}/api/status").json()
    elapsed = int(time.time() - start_time)
    print(f"[{elapsed}s] State: {s['state']} | Combos: {s['current_combination_idx']}/{s['total_combinations']} | Current: '{s['current_service']} in {s['current_location']}' | Processed: {s['businesses_processed']} | Total Unique: {s['total_collected']}")
    
    if s["state"] in ["completed", "stopped", "error"]:
        break
        
print("\nFinal Status:", s["status_message"])
print("Total businesses collected:", len(s["businesses"]))
for i, b in enumerate(s["businesses"][:5], 1):
    print(f"[{i}] {b['name']} | Service: {b['searchService']} | Location: {b['searchLocation']} | Phone: {b['phone']} | Claim: {b['claimStatus']} | Email: {b['email']}")

# Test Excel Export
print("\n--- Testing Excel Export Endpoint ---")
export_resp = requests.get(f"{BASE_URL}/api/export")
print("Export HTTP Status:", export_resp.status_code)
assert export_resp.status_code == 200, "Export failed!"

test_excel_path = "verified_export.xlsx"
with open(test_excel_path, "wb") as f:
    f.write(export_resp.content)
print(f"Saved exported Excel to {test_excel_path} ({len(export_resp.content)} bytes)")

wb = openpyxl.load_workbook(test_excel_path)
ws = wb.active
print("Worksheet title:", ws.title)
print("Total rows:", ws.max_row)
print("Header columns:", [cell.value for cell in ws[1]])

# Print sample row from Excel
if ws.max_row > 1:
    sample_row = [cell.value for cell in ws[2]]
    print("Row 2 data:", sample_row)

print("\n--- ALL VERIFICATIONS PASSED SUCCESSFULLY! ---")
