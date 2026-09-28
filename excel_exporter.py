import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Business Name", 32),
    ("Category", 26),
    ("Website", 30),
    ("Email", 28),
    ("Phone", 20),
    ("Address", 38),
    ("City", 20),
    ("State", 16),
    ("Country", 14),
    ("Google Maps URL", 35),
    ("Rating", 12),
    ("Reviews", 14),
    ("Claim Status", 18),
    ("Search Service", 24),
    ("Search Location", 22),
]

def export_to_excel(businesses, output_filepath):
    """
    Exports a list of business dicts to a beautifully styled Excel spreadsheet.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Google Maps Businesses"
    
    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True
    
    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Border styles
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Write Header
    ws.row_dimensions[1].height = 28
    for col_idx, (col_name, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
        
    # Styling for data rows
    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10, color="0F172A")
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    # Status badges styling
    claim_avail_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    claim_avail_font = Font(name="Segoe UI", size=10, bold=True, color="166534")
    
    claim_not_avail_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    claim_not_avail_font = Font(name="Segoe UI", size=10, color="64748B")

    claim_unknown_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    claim_unknown_font = Font(name="Segoe UI", size=10, color="92400E")

    for row_idx, b in enumerate(businesses, start=2):
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)
        default_fill = row_alt_fill if is_even else None
        
        row_data = [
            b.get("name", "Not Available") or "Not Available",
            b.get("category", "Not Available") or "Not Available",
            b.get("website", "Not Available") or "Not Available",
            b.get("email", "Not Available") or "Not Available",
            b.get("phone", "Not Available") or "Not Available",
            b.get("address", "Not Available") or "Not Available",
            b.get("city", "Not Available") or "Not Available",
            b.get("state", "Not Available") or "Not Available",
            b.get("country", "Not Available") or "Not Available",
            b.get("url", "Not Available") or "Not Available",
            b.get("rating", "Not Available") or "Not Available",
            b.get("reviews", "Not Available") or "Not Available",
            b.get("claimStatus", "Not Available") or "Not Available",
            b.get("searchService", "Not Available") or "Not Available",
            b.get("searchLocation", "Not Available") or "Not Available",
        ]
        
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if default_fill:
                cell.fill = default_fill
                
            # Formatting specifics
            header_name = COLUMNS[col_idx - 1][0]
            if header_name in ["Rating", "Reviews", "Phone", "City", "State", "Country"]:
                cell.alignment = center_align
                cell.value = val
            elif header_name == "Claim Status":
                cell.alignment = center_align
                cell.value = val
                if val == "Available":
                    cell.fill = claim_avail_fill
                    cell.font = claim_avail_font
                elif val == "Not Available":
                    cell.fill = claim_not_avail_fill
                    cell.font = claim_not_avail_font
                elif val == "Unknown":
                    cell.fill = claim_unknown_fill
                    cell.font = claim_unknown_font
            elif header_name in ["Website", "Google Maps URL"]:
                cell.alignment = left_align
                cell.value = val
                if val and val.startswith("http"):
                    cell.font = link_font
                    cell.hyperlink = val
            else:
                cell.alignment = left_align
                cell.value = val
                
    # Freeze Header Row
    ws.freeze_panes = "A2"
    
    # Auto Filter
    if len(businesses) > 0:
        max_col_letter = get_column_letter(len(COLUMNS))
        ws.auto_filter.ref = f"A1:{max_col_letter}{len(businesses) + 1}"
    else:
        max_col_letter = get_column_letter(len(COLUMNS))
        ws.auto_filter.ref = f"A1:{max_col_letter}1"
        
    # Auto-fit Column Widths (with min configured padding)
    for col_idx, (_, default_width) in enumerate(COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(default_width, min(max_len + 4, 60))
        
    os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
    wb.save(output_filepath)
    return output_filepath

if __name__ == '__main__':
    # Test sample export
    sample = [
        {
            "name": "D.B. Container Service",
            "category": "Dumpster rental service",
            "website": "https://www.dbcontainers.com/",
            "email": "Not Available",
            "phone": "(718) 257-2300",
            "address": "129 Louisiana Ave, Brooklyn, NY 11207",
            "city": "Brooklyn",
            "state": "NY",
            "country": "USA",
            "url": "https://maps.google.com/?cid=123",
            "rating": "4.8",
            "reviews": "121",
            "claimStatus": "Available",
            "searchService": "Dumpster Rental",
            "searchLocation": "New York"
        }
    ]
    path = export_to_excel(sample, "test_output.xlsx")
    print(f"Successfully exported test file to: {path}")
